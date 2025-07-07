import os
import time
import datetime
import base64
import io

from io import StringIO, BytesIO
from collections import OrderedDict, Counter

import xlsxwriter
from flask import (
    Flask, render_template, redirect, url_for, flash, request,
    jsonify, send_file, send_from_directory, abort
)
from flask_migrate import Migrate
from flask_login import LoginManager, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from sqlalchemy import func, case

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl import Workbook
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart import BarChart, PieChart, Reference

from config import Config
from models import db, User, ImageRequest
from forms import LoginForm, RegisterForm
import image_pipeline as ip

app = Flask(__name__)
app.config.from_object(Config)
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 0

db.init_app(app)
Migrate(app, db)

login_manager = LoginManager(app)
login_manager.login_view = 'login'

@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

UPLOAD_FOLDER = "uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

@app.route('/')
@login_required
def index():
    return render_template('index.html')

@app.route('/register', methods=['GET', 'POST'])
def register():
    if current_user.is_authenticated:
        return redirect(url_for('index'))
    form = RegisterForm()
    if form.validate_on_submit():
        if User.query.filter_by(username=form.username.data).first():
            flash('Пользователь уже существует', 'danger')
        else:
            u = User(
                username=form.username.data,
                password_hash=generate_password_hash(form.password.data)
            )
            db.session.add(u)
            db.session.commit()
            flash('Регистрация прошла успешно, войдите', 'success')
            return redirect(url_for('login'))
    else:
        for errs in form.errors.values():
            for e in errs:
                flash(e, 'danger')
    return render_template('register.html', form=form)

@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('index'))
    form = LoginForm()
    if form.validate_on_submit():
        u = User.query.filter_by(username=form.username.data).first()
        if u and check_password_hash(u.password_hash, form.password.data):
            login_user(u)
            return redirect(url_for('index'))
        else:
            flash('Неверное имя или пароль', 'danger')
    else:
        for errs in form.errors.values():
            for e in errs:
                flash(e, 'danger')
    return render_template('login.html', form=form)

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))

@app.route('/history')
@login_required
def history_page():
    return render_template('history.html')

@app.route('/api/classify', methods=['POST'])
@login_required
def api_classify():
    if 'file' not in request.files:
        return jsonify(error='No file'), 400
    f = request.files['file']
    if not ip.allowed_file(f.filename):
        return jsonify(error='Invalid type'), 400

    fname = secure_filename(f.filename)
    os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
    path = os.path.join(app.config['UPLOAD_FOLDER'], fname)
    f.save(path)

    with open(path, 'rb') as img_f:
        img_blob = img_f.read()

    start_ts = time.perf_counter()
    try:
        verdict, prob = ip.classify_image(path)
    except Exception as e:
        verdict, prob = str(e), None
    processing_time = time.perf_counter() - start_ts

    rec = ImageRequest(
        filename        = fname,
        verdict         = verdict,
        probability     = prob,
        processing_time = processing_time,
        image_data      = img_blob,
        user_id         = current_user.id
    )
    db.session.add(rec)
    db.session.commit()

    try:
        os.remove(path)
    except OSError:
        pass

    img_b64 = base64.b64encode(img_blob).decode('utf-8')
    return jsonify(
        id              = rec.id,
        filename        = fname,
        verdict         = verdict,
        probability     = prob,
        processing_time = processing_time,
        image_data      = img_b64,
        timestamp       = rec.timestamp.isoformat() + 'Z'
    ), 200


@app.route('/api/history', methods=['GET'])
@login_required
def api_history():
    page = int(request.args.get('page', 1))
    per_page = int(request.args.get('per_page', 100))
    limit = request.args.get('limit', None)
    meta_only = request.args.get('meta', 'false').lower() == 'true'
    status_filter = request.args.get('status', 'all')

    base_q = ImageRequest.query.filter_by(user_id=current_user.id).order_by(ImageRequest.timestamp.desc())
    if status_filter != 'all':
        base_q = base_q.filter(ImageRequest.verdict == status_filter)

    if limit is not None:
        recs = base_q.limit(int(limit)).all()
    else:
        recs = base_q.offset((page-1)*per_page).limit(per_page).all()

    history = []
    for r in recs:
        entry = {
            'id': r.id,
            'timestamp': r.timestamp.isoformat() + 'Z',
            'filename': r.filename,
            'verdict': r.verdict,
            'probability': r.probability
        }
        if not meta_only:
            entry['image_data'] = base64.b64encode(r.image_data).decode('utf-8')
        history.append(entry)

    resp = {'history': history}
    if meta_only:
        resp['total'] = base_q.count()
    return jsonify(resp), 200

@app.route('/api/history/<int:rec_id>/image')
@login_required
def api_history_image(rec_id):
    rec = ImageRequest.query.get_or_404(rec_id)
    if rec.user_id != current_user.id:
        abort(403)
    return send_file(BytesIO(rec.image_data), mimetype='image/png')

@app.route('/api/history', methods=['DELETE'])
@login_required
def api_delete_history():
    ImageRequest.query.filter_by(user_id=current_user.id).delete()
    db.session.commit()
    return '', 204

@app.route('/uploads/<path:filename>')
@login_required
def uploaded_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)

@app.route('/stats')
@login_required
def stats_page():
    return render_template('stats.html')

@app.route('/api/stats')
@login_required
def api_stats():
    date_from = request.args.get('from')
    date_to   = request.args.get('to')

    dt_from = None
    dt_to   = None
    if date_from:
        d = datetime.date.fromisoformat(date_from)
        dt_from = datetime.datetime.combine(d, datetime.time.min)
    if date_to:
        d = datetime.date.fromisoformat(date_to)
        dt_to = datetime.datetime.combine(d + datetime.timedelta(days=1), datetime.time.min)

    q = ImageRequest.query.filter_by(user_id=current_user.id)
    if dt_from:
        q = q.filter(ImageRequest.timestamp >= dt_from)
    if dt_to:
        q = q.filter(ImageRequest.timestamp < dt_to)

    total = q.with_entities(func.count()).scalar() or 0

    counts_row = db.session.query(
        func.sum(case((ImageRequest.verdict=='good',1), else_=0)).label('good'),
        func.sum(case((ImageRequest.verdict=='defective',1), else_=0)).label('defective'),
        func.sum(case((ImageRequest.verdict=='external',1), else_=0)).label('external'),
        func.sum(case((ImageRequest.verdict=='error',1), else_=0)).label('error'),
    ).filter_by(user_id=current_user.id)
    if dt_from:
        counts_row = counts_row.filter(ImageRequest.timestamp >= dt_from)
    if dt_to:
        counts_row = counts_row.filter(ImageRequest.timestamp < dt_to)
    counts_row = counts_row.one()

    counts = OrderedDict([
        ('good',      counts_row.good      or 0),
        ('defective', counts_row.defective or 0),
        ('external',  counts_row.external  or 0),
        ('error',     counts_row.error     or 0),
    ])
    percents = {k: (counts[k]/total*100 if total else 0) for k in counts}

    if date_from and date_to:
        d0 = datetime.date.fromisoformat(date_from)
        d1 = datetime.date.fromisoformat(date_to)
        selected_days = (d1 - d0).days + 1
    elif date_from or date_to:
        selected_days = 1
    else:
        min_ts, max_ts = db.session.query(
            func.min(ImageRequest.timestamp),
            func.max(ImageRequest.timestamp)
        ).filter_by(user_id=current_user.id).one()
        if min_ts and max_ts and max_ts >= min_ts:
            selected_days = (max_ts.date() - min_ts.date()).days + 1
        else:
            selected_days = 1

    avg_per_day = (total/selected_days) if selected_days > 0 else 0

    recs_q = ImageRequest.query.filter_by(user_id=current_user.id)
    if dt_from:
        recs_q = recs_q.filter(ImageRequest.timestamp >= dt_from)
    if dt_to:
        recs_q = recs_q.filter(ImageRequest.timestamp < dt_to)
    recs = recs_q.order_by(ImageRequest.timestamp).all()
    records = [{'ts': r.timestamp.isoformat(), 'verdict': r.verdict} for r in recs]

    t_q = db.session.query(
        func.min(ImageRequest.processing_time),
        func.max(ImageRequest.processing_time),
        func.avg(ImageRequest.processing_time)
    ).filter_by(user_id=current_user.id)
    if dt_from:
        t_q = t_q.filter(ImageRequest.timestamp >= dt_from)
    if dt_to:
        t_q = t_q.filter(ImageRequest.timestamp < dt_to)
    min_t, max_t, avg_t = t_q.one()

    return jsonify({
      "total": total,
      "counts": counts,
      "percents": percents,
      "metrics": {
        "selected_days": selected_days,
        "avg_per_day":    avg_per_day,
        "min_time":       float(min_t or 0),
        "max_time":       float(max_t or 0),
        "avg_time":       float(avg_t or 0)
      },
      "records": records
    })


def format_num(v):
    if v is None:
        return "0"
    if v % 1 == 0:
        return str(int(v))
    return str(round(v, 2))


def format_time(sec):
    if sec is None:
        return "0 мс"
    ms = sec * 1000
    if ms < 1000:
        return f"{int(ms)} мс"
    return f"{round(ms/1000, 2)} с"


@app.route('/api/stats/export.xlsx')
@login_required
def export_stats_xlsx():
    data = api_stats().get_json()
    recs = data['records']

    date_from = request.args.get('from')
    date_to   = request.args.get('to')
    if date_from and date_to:
        start_date = datetime.date.fromisoformat(date_from)
        end_date   = datetime.date.fromisoformat(date_to)
    elif recs:
        start_date = datetime.date.fromisoformat(recs[0]['ts'][:10])
        end_date   = datetime.date.fromisoformat(recs[-1]['ts'][:10])
    else:
        start_date = end_date = datetime.date.today()
    single_day = (start_date == end_date)
    period_label = (
        start_date.isoformat()
        if single_day
        else f"{start_date.isoformat()} – {end_date.isoformat()}"
    )

    output = io.BytesIO()
    workbook = xlsxwriter.Workbook(output, {'in_memory': True})

    fmt_center = workbook.add_format({'align': 'center', 'valign': 'vcenter'})
    fmt_pct    = workbook.add_format({'num_format': '0.0%', 'align': 'center', 'valign': 'vcenter'})
    colors = {
        'good':      '#2e7d32',
        'defective': '#2852c6',
        'external':  '#b910bc',
        'error':     '#ef6c00',
    }

    ws1 = workbook.add_worksheet("Сводка")
    ws1.set_column('A:C', 25, fmt_center)

    ws1.write(0, 0, "Период", fmt_center)
    ws1.write(0, 1, period_label, fmt_center)

    ws1.write_row(2, 0, ["Показатель", "Значение", "% (для категорий)"], fmt_center)
    ws1.write_row(3, 0, ["Всего запросов", data['total'], None], fmt_center)

    row = 4
    for key, label in [
        ('good',      'Пригодные'),
        ('defective', 'Непригодные'),
        ('external',  'Сторонние'),
        ('error',     'Отклонённые'),
    ]:
        ws1.write_row(row, 0,
                      [label,
                       data['counts'][key],
                       data['percents'][key] / 100.0],
                      fmt_center if key != None else fmt_pct)
        ws1.write(row, 2, data['percents'][key] / 100.0, fmt_pct)
        row += 1
    row += 1

    extras = [
        ("Среднее запросов в день", round(data['metrics']['avg_per_day'], 2)),
        ("Мин. время обработки", format_time(data['metrics']['min_time'])),
        ("Макс. время обработки", format_time(data['metrics']['max_time'])),
        ("Ср. время обработки", format_time(data['metrics']['avg_time'])),
    ]
    for txt, val in extras:
        ws1.write(row, 0, txt, fmt_center)
        ws1.write(row, 1, val, fmt_center)
        row += 1

    ws2 = workbook.add_worksheet("Гистограмма")
    ws2.set_column('A:E', 17, fmt_center)

    headers = ["Метка", "Пригодные", "Непригодные", "Сторонние", "Отклонённые"]
    ws2.write_row(0, 0, headers, fmt_center)

    if single_day:
        labels = [f"{h:02d}:00" for h in range(24)]
    else:
        labels = []
        cur = start_date
        while cur <= end_date:
            labels.append(cur.isoformat())
            cur += datetime.timedelta(days=1)

    from collections import OrderedDict
    buckets = OrderedDict((lbl, {'good':0,'defective':0,'external':0,'error':0})
                          for lbl in labels)
    for r in recs:
        ts = datetime.datetime.fromisoformat(r['ts'])
        lbl = (ts.strftime("%H:00") if single_day else ts.date().isoformat())
        buckets[lbl][r['verdict']] += 1

    for i, (lbl, cnts) in enumerate(buckets.items(), start=1):
        ws2.write(i, 0, lbl, fmt_center)
        ws2.write(i, 1, cnts['good'], fmt_center)
        ws2.write(i, 2, cnts['defective'], fmt_center)
        ws2.write(i, 3, cnts['external'], fmt_center)
        ws2.write(i, 4, cnts['error'], fmt_center)

    chart = workbook.add_chart({'type': 'column', 'subtype': 'stacked'})
    for col, key in enumerate(['good','defective','external','error'], start=1):
        chart.add_series({
            'name':       headers[col],
            'categories': ['Гистограмма', 1, 0, len(labels), 0],
            'values':     ['Гистограмма', 1, col, len(labels), col],
            'fill':       {'color': colors[key]},
        })
    chart.set_x_axis({'num_font': {'rotation': -45}})
    chart.set_legend({'position': 'bottom'})
    chart.set_size({'width': 720, 'height': 360})
    ws2.insert_chart('G2', chart)

    ws3 = workbook.add_worksheet("Диаграмма")
    ws3.set_column('A:B', 17, fmt_center)

    ws3.write_row(0, 0, ["Категория", "Значение"], fmt_center)
    pie_data = [(key, label, data['counts'][key])
                for key, label in [
                    ('good','Пригодные'),
                    ('defective','Непригодные'),
                    ('external','Сторонние'),
                    ('error','Отклонённые'),
                ] if data['counts'][key] > 0]
    for i, (_, label, val) in enumerate(pie_data, start=1):
        ws3.write(i, 0, label, fmt_center)
        ws3.write(i, 1, val,   fmt_center)

    pie = workbook.add_chart({'type': 'pie'})
    pie.add_series({
        'categories': ['Диаграмма', 1, 0, len(pie_data), 0],
        'values':     ['Диаграмма', 1, 1, len(pie_data), 1],
        'data_labels': {'percentage': True, 'position': 'outside_end'},
        'points':     [{'fill': {'color': colors[k]}} for k, _, _ in pie_data],
    })
    pie.set_legend({'position': 'bottom'})
    pie.set_size({'width': 360, 'height': 360})
    ws3.insert_chart('D2', pie)

    workbook.close()
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="statistics.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


if __name__ == '__main__':
    app.run(debug=True)
